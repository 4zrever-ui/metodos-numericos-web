"""Compara las hojas recalculadas con method_all; hallazgos H12 a H15."""

import math
import numbers
import re
from pathlib import Path

import pytest
import sympy

formulas = pytest.importorskip("formulas")

import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

import backend.main as api
from backend.excel.excel_generator import ALL_METHODS, SHEET_NAMES, generate_all


CASES = (
    ("x^2-2", 1.0, 2.0, 0.0, 2.0, "(x+2)/(x+1)"),
    ("x^3-2*x-5", 2.0, 3.0, 2.0, 3.0, "(2*x+5)^(1/3)"),
)
FPP_METHODS = {
    "newton_modificado", "newton_2do_orden", "chebyshev",
    "halley", "super_halley", "ostrowsky",
}
MANUAL_SIZE_METHODS = {"biseccion", "regula_falsi", "punto_fijo", "aitken"}
API_KEYS = {
    "newton_raphson": "newton",
    "newton_2do_orden": "newton_segundo_orden",
}
COLUMNS = {
    "biseccion": (3, 10), "regula_falsi": (3, 10),
    "punto_fijo": (3, 5), "aitken": (2, 4), "steffensen": (2, 4),
    "newton_raphson": (5, 7), "von_mises": (5, 7), "secante": (6, 8),
}


def _payload(case, mode):
    equation, x0, x1, a, b, gx = CASES[case]
    data = {"equation": equation}
    if mode == "B":
        data.update(x0=x0, x1=x1, a=a, b=b, gx=gx)
    elif mode == "C":
        data["x0"] = x0
    return data


def _overrides(payload):
    return {
        key: params for key in ALL_METHODS
        if (params := api._excel_params(key, payload))
    } or None


def _backend(payload):
    """Observa filas de los runners reales, restaurando la lista al terminar."""
    captured = {}

    def wrap(key, runner):
        def observe(eq, params, data):
            result = runner(eq, params, data)
            captured[key] = result
            return result
        return observe

    original = api.ALL_METHOD_RUNNERS
    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(api, "ALL_METHOD_RUNNERS", [
            (key, wrap(key, runner)) for key, runner in original
        ])
        response = api.method_all(payload)
    assert set(captured) == {row["method"] for row in response["results"]}
    for row in response["results"]:
        assert captured[row["method"]].converged == row["converged"]
    return captured


def _solution_cells(solution):
    """Misma lectura por hoja/coordenada que el script externo de auditoria."""
    values = {}
    for node, result in solution.items():
        key = str(node)
        if "!" not in key or "]" not in key:
            continue
        prefix, address = key.rsplit("!", 1)
        sheet = prefix.split("]", 1)[1].strip("'").replace("''", "'").upper()
        address = address.replace("$", "")
        if not re.fullmatch(r"[A-Z]+[0-9]+(?::[A-Z]+[0-9]+)?", address):
            continue
        grid = np.asarray(result.value, dtype=object)
        if grid.ndim == 0:
            grid = grid.reshape((1, 1))
        elif grid.ndim == 1:
            grid = grid.reshape((1, -1))
        min_c, min_r, max_c, max_r = range_boundaries(address)
        assert grid.shape == (max_r - min_r + 1, max_c - min_c + 1)
        for ri in range(grid.shape[0]):
            for ci in range(grid.shape[1]):
                coord = f"{get_column_letter(min_c + ci)}{min_r + ri}"
                values[sheet, coord] = grid[ri, ci]
    return values


def _numeric(value):
    return (
        isinstance(value, numbers.Real)
        and not isinstance(value, (bool, np.bool_))
        and math.isfinite(value)
    )


def _sheet_result(ws, method, values):
    xcol, convcol = COLUMNS.get(method, (6, 8))

    def value(cell):
        if cell.data_type == "f":
            return values[ws.title.upper(), cell.coordinate]
        return cell.value

    rows, verdicts = {}, []
    for cells in ws:
        k = cells[0].value
        if not isinstance(k, int) or isinstance(k, bool):
            continue
        # Secante: F2 (k=0) no tiene x_new; se compara solo k >= 1.
        if method == "secante" and k == 0:
            continue
        r = cells[0].row
        x_new = value(ws.cell(r, xcol))
        if _numeric(x_new):
            rows[k] = float(x_new)
        verdicts.append(value(ws.cell(r, convcol)))
    return rows, any(str(v).strip().upper() == "SI" for v in verdicts)


@pytest.fixture(scope="module")
def recalculated_books(tmp_path_factory):
    """Genera y recalcula una vez cada uno de los cuatro libros A/B."""
    directory = tmp_path_factory.mktemp("excel_vs_backend")
    repo = Path(__file__).resolve().parents[2]
    assert directory != repo and repo not in directory.parents
    books = {}
    for case in range(len(CASES)):
        for mode in ("A", "B"):
            payload = _payload(case, mode)
            backend = _backend(payload)
            path = directory / f"case_{case + 1}_{mode}.xlsx"
            path.write_bytes(generate_all(
                payload["equation"],
                params_per_method=None if mode == "A" else _overrides(payload),
                eq_label=payload["equation"],
            ))
            model = formulas.ExcelModel().loads(str(path)).finish()
            values = _solution_cells(model.calculate())
            workbook = openpyxl.load_workbook(path, data_only=False)
            try:
                excel = {
                    method: _sheet_result(workbook[SHEET_NAMES[method]], method, values)
                    for method in ALL_METHODS
                }
            finally:
                workbook.close()
            books[case, mode] = (backend, excel)
    return books


def _comparisons():
    for case in range(len(CASES)):
        for mode in ("A", "B"):
            for method in ALL_METHODS:
                marks = []
                if case == 1 and method in FPP_METHODS:
                    reason = "H14: Excel fija f''(x0); backend evalua f''(xk)"
                elif mode == "B" and method in MANUAL_SIZE_METHODS:
                    reason = "H13: Excel dimensiona con parametros automaticos"
                elif case == 1 and mode == "A" and method == "steffensen":
                    reason = "H15: J10 vacia por delta2 cero; B3 queda vacia"
                else:
                    reason = None
                if reason:
                    marks.append(pytest.mark.xfail(
                        strict=True, reason=reason, raises=AssertionError,
                    ))
                yield pytest.param(
                    case, mode, method, marks=marks,
                    id=f"case{case + 1}-{mode}-{method}",
                )


@pytest.mark.parametrize("case,mode,method", list(_comparisons()))
def test_excel_matches_backend(recalculated_books, case, mode, method):
    backend, excel = recalculated_books[case, mode]
    result = backend[API_KEYS.get(method, method)]
    expected = {
        row.k: row.x_new for row in result.iterations
        if method != "secante" or row.k >= 1
    }
    actual, converged = excel[method]
    assert len(actual) == len(expected), "Numero de filas numericas Excel/backend"
    assert actual.keys() == expected.keys(), "Indices k Excel/backend"
    for k, x_new in expected.items():
        assert actual[k] == pytest.approx(x_new, abs=1e-12, rel=0), f"x_new en k={k}"
    assert converged == result.converged, "Veredicto de convergencia Excel/backend"


@pytest.mark.parametrize("case", [
    pytest.param(0, id="case1-C"),
    pytest.param(1, id="case2-C"),
])
@pytest.mark.xfail(
    strict=True,
    reason="H12: overrides con x0 sin gx no completan g_str",
    raises=sympy.SympifyError,
)
def test_excel_x0_without_gx_does_not_raise(case):
    payload = _payload(case, "C")
    contents = generate_all(
        payload["equation"],
        params_per_method=_overrides(payload),
        eq_label=payload["equation"],
    )
    assert contents

