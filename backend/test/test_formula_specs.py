"""
Fija la coincidencia entre `excel/formula_specs.py` y las plantillas reales.

Por qué este test existe
------------------------
`formula_specs.py` duplica, como dato, las fórmulas que `excel_templates.py`
escribe en la hoja. Duplicar invita a divergir: alguien cambia la plantilla y la
especificación que ve el alumno en Practicar se queda mintiendo.

Aquí se construye la hoja DE VERDAD con la plantilla real, en memoria, y se
comprueba celda a celda que dice lo mismo que la especificación. Si la plantilla
cambia, esto se rompe.

De paso, es el PRIMER test que toca `excel_templates.py`: hasta ahora ese archivo
—1091 líneas— no lo importaba ninguna prueba de la suite (hallazgo H3).
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook

from backend.core.auto_params import generate_params
from backend.core.equation_parser import parse_equation
from backend.excel.excel_templates import BiseccionTemplate, NewtonRaphsonTemplate
from backend.excel.formula_specs import spec_para, spec_serializable

ECUACION = "x**3 - 2*x - 5"


@pytest.fixture(scope="module")
def contexto():
    eq = parse_equation(ECUACION)
    p = generate_params(eq)
    return {
        "fx": str(eq.f_sympy),
        "fpx": str(eq.fp_sympy),
        "x0": p.x0,
        "a0": p.a,
        "b0": p.b,
    }


@pytest.fixture(scope="module")
def hojas(contexto):
    """Las dos hojas construidas por las plantillas reales, en memoria."""
    wb = Workbook()

    ws_newton = wb.active
    ws_newton.title = "Newton"
    NewtonRaphsonTemplate().build(
        ws_newton, contexto["fx"], {"x0": contexto["x0"]},
        n_iter=6, eq_label=ECUACION, fpx_str=contexto["fpx"],
    )

    ws_bis = wb.create_sheet("Biseccion")
    BiseccionTemplate().build(
        ws_bis, contexto["fx"], {"a0": contexto["a0"], "b0": contexto["b0"]},
        n_iter=6, eq_label=ECUACION,
    )

    return {"newton": ws_newton, "biseccion": ws_bis}


def desenvolver(valor: str, conv_prev: str, pivot_prev: str) -> str:
    """
    Quita el envoltorio `_freeze` y devuelve la fórmula limpia, sin el '=' inicial.

    Las tres formas que produce `_freeze` (ver su docstring):
        k = 0   la fórmula tal cual
        k = 1   IF({conv}="SI","",{formula})
        k >= 2  IF(OR({conv}="SI",{pivot}=""),"",{formula})
    """
    cuerpo = valor[1:] if valor.startswith("=") else valor

    for prefijo in (
        f'IF({conv_prev}="SI","",',
        f'IF(OR({conv_prev}="SI",{pivot_prev}=""),"",',
    ):
        if cuerpo.startswith(prefijo):
            assert cuerpo.endswith(")"), f"envoltorio sin cerrar: {cuerpo!r}"
            return cuerpo[len(prefijo):-1]

    return cuerpo


# Columna de convergencia y columna pivote de cada método, que son las que
# `_freeze` mira en la fila anterior.
REFS_FREEZE = {"newton": ("G", "B"), "biseccion": ("J", "C")}


def celda(ws, columna: str, fila: int):
    return ws[f"{columna}{fila}"].value


@pytest.mark.parametrize("metodo", ["newton", "biseccion"])
@pytest.mark.parametrize("fila", [3, 4, 5])
def test_la_especificacion_coincide_con_la_plantilla(hojas, contexto, metodo, fila):
    """
    En las filas de régimen (k >= 1) TODAS las columnas son fórmulas, y cada una
    debe coincidir con la especificación una vez quitado el envoltorio.
    """
    spec = spec_para(metodo, contexto["fx"], contexto["fpx"])
    ws = hojas[metodo]
    col_conv, col_pivot = REFS_FREEZE[metodo]
    conv_prev, pivot_prev = f"{col_conv}{fila-1}", f"{col_pivot}{fila-1}"

    for columna in spec["columnas"]:
        letra = columna["columna"]
        if letra == "A":
            continue  # el índice k es un literal entero, no una fórmula

        real = celda(ws, letra, fila)
        assert isinstance(real, str) and real.startswith("="), (
            f"{metodo} {letra}{fila}: se esperaba una fórmula, hay {real!r}"
        )

        esperada = columna["formula"](fila)
        obtenida = desenvolver(real, conv_prev, pivot_prev)
        assert obtenida == esperada, (
            f"{metodo} {letra}{fila} no coincide.\n"
            f"  plantilla:      {obtenida}\n"
            f"  especificación: {esperada}"
        )


def test_newton_fila_k0_ya_trae_las_formulas_limpias(hojas, contexto):
    """
    En k=0 `_freeze` es la identidad, así que la fila 2 de Newton contiene las
    fórmulas SIN envoltorio. Es la evidencia del sondeo D2.
    """
    spec = spec_para("newton", contexto["fx"], contexto["fpx"])
    ws = hojas["newton"]

    assert celda(ws, "B", 2) == pytest.approx(contexto["x0"]), "en k=0, B es el x0 literal"

    for letra in ("C", "D", "E", "F", "G"):
        columna = next(c for c in spec["columnas"] if c["columna"] == letra)
        real = celda(ws, letra, 2)
        assert real.startswith("="), f"{letra}2 deberia ser formula, es {real!r}"
        assert "IF(OR(" not in real, f"{letra}2 no deberia llevar envoltorio en k=0"
        assert real[1:] == columna["formula"](2), f"{letra}2 no coincide con la especificacion"


def test_biseccion_en_k0_NO_es_simetrica_con_newton(hojas, contexto):
    """
    El caso que el sondeo D2 obligó a mirar aparte, y que no se puede deducir
    por simetría con Newton: en la fila k=0 de Bisección
      - B y D son los LITERALES del intervalo inicial, no fórmulas;
      - I y J (error y convergencia) están VACÍAS.
    """
    ws = hojas["biseccion"]

    assert celda(ws, "B", 2) == pytest.approx(contexto["a0"]), "a0 literal en k=0"
    assert celda(ws, "D", 2) == pytest.approx(contexto["b0"]), "b0 literal en k=0"

    for letra in ("I", "J"):
        valor = celda(ws, letra, 2)
        assert valor in ("", None), f"{letra}2 deberia estar vacia en k=0, hay {valor!r}"

    # Las demás sí son fórmulas ya en k=0.
    for letra in ("C", "E", "F", "G", "H"):
        assert str(celda(ws, letra, 2)).startswith("="), f"{letra}2 deberia ser formula"


def test_biseccion_actualiza_el_intervalo_desde_k1_y_envuelto(hojas, contexto):
    """
    Las dos fórmulas más interesantes de Bisección —las que deciden con qué mitad
    se sigue— NO existen en k=0: aparecen en k=1 y ya envueltas en `_freeze`.
    Se verifican explícitamente, que es lo que pidió el director.
    """
    spec = spec_para("biseccion", contexto["fx"])
    ws = hojas["biseccion"]

    for letra in ("B", "D"):
        columna = next(c for c in spec["columnas"] if c["columna"] == letra)
        real = celda(ws, letra, 3)

        assert real.startswith('=IF(J2="SI","",'), (
            f"{letra}3 deberia venir envuelta por _freeze (k=1), es {real!r}"
        )
        obtenida = desenvolver(real, "J2", "C2")
        assert obtenida == columna["formula"](3)
        assert obtenida.startswith("IF(E2*F2<0,"), (
            f"{letra}3 deberia decidir por el signo de f(a)*f(c), es {obtenida!r}"
        )


def test_serializable_no_lleva_lambdas_y_trae_el_igual(contexto):
    """Lo que viaja al frontend: fórmulas ya resueltas, con '=' y sin funciones."""
    import json

    datos = spec_serializable("newton", contexto["fx"], contexto["fpx"], fila=3)
    json.dumps(datos)   # revienta si quedara una lambda

    assert datos["k"] == 1
    assert len(datos["columnas"]) == 7
    for c in datos["columnas"]:
        assert c["formula"].startswith("="), c
        assert c["explicacion"], "cada columna necesita su explicacion para el alumno"

    assert spec_serializable("halley", contexto["fx"]) is None, "metodo no cubierto -> None"
